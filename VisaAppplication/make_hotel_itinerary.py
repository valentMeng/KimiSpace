# -*- coding: utf-8 -*-
"""按参考图片格式生成酒店行程单（Italy-France 2026-11-20 ~ 2026-12-02）"""
import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

HEADERS = ["Day", "Date", "Country", "City", "Trip Itinerary", "Touring spot", "Transportation", "accommodation"]

ROWS = [
    # Day, Date, Country, City, Trip Itinerary, Touring spot, Transportation, accommodation
    (1, "2026/11/20\nFriday", "China - France",
     "Shanghai → Paris\n上海 → 巴黎",
     "Hangzhou → Shanghai\n杭州 → 上海\nDeparture from Shanghai\n出发",
     "Hangzhou East Railway Station\n杭州东站\nShanghai Hongqiao Airport T2\n上海虹桥机场T2",
     "Flight: CA1566 22:30\n（上海虹桥T2 → 北京）\nFlight: CA875 02:45\n（北京 → 巴黎戴高乐T1）",
     "On board flight\n（宿飞机上）"),
    (2, "2026/11/21\nSaturday", "France",
     "Paris\n巴黎",
     "Arrive at CDG T1 07:00 → transfer to hotel & check in\n抵达戴高乐T1 → 前往酒店办理入住\nLe Marais stroll & lunch\n玛黑区漫步、午餐",
     "Musée d'Orsay\n奥赛博物馆\nEiffel Tower\n埃菲尔铁塔\nSeine river cruise (19:00)\n塞纳河游船（19:00）",
     "Flight + Metro\n（航班 + 地铁）",
     "Kopster Hotel Paris Porte de Versailles\nAddress: 1 Boulevard Victor,\n75015 Paris\nPhone: +33 1 76 49 05 00"),
    (3, "2026/11/22\nSunday", "France",
     "Paris\n巴黎",
     "Hotel → Louvre (morning, pre-booked)\n酒店 → 卢浮宫（上午，已预订）\nWalk along the Seine\n沿塞纳河步行",
     "Louvre Museum\n卢浮宫\nTuileries Garden\n杜乐丽花园\nPlace de la Concorde\n协和广场\nHôtel de la Marine\n海军府\nPalais Garnier (evening)\n巴黎歌剧院（晚）",
     "Metro / Walk\n（地铁 / 步行）",
     "Kopster Hotel Paris Porte de Versailles\nAddress: 1 Boulevard Victor,\n75015 Paris\nPhone: +33 1 76 49 05 00"),
    (4, "2026/11/23\nMonday", "France",
     "Paris\n巴黎",
     "Hotel → Luxembourg Gardens (morning)\n酒店 → 卢森堡公园（上午）\nLeft Bank walk & bookstores\n左岸漫步、书店",
     "Luxembourg Gardens\n卢森堡公园\nShakespeare and Company\n莎士比亚书店\nNotre-Dame de Paris\n巴黎圣母院\nMontmartre & Sacré-Cœur (evening)\n蒙马特 & 圣心大教堂（晚）",
     "Metro\n（地铁）",
     "Holiday Inn Paris – Gare de Lyon Bastille\nAddress: 11 Rue de Lyon,\n75012 Paris\nPhone: +33 1 43 42 36 10"),
    (5, "2026/11/24\nTuesday", "France",
     "Paris → Nice\n巴黎 → 尼斯",
     "OUIGO train 06:03 Paris Gare de Lyon → Nice Ville 12:38\nOUIGO高铁 06:03 巴黎里昂车站 → 尼斯 12:38\nHotel check-in\n酒店入住\nVieux Nice & Promenade des Anglais\n尼斯老城 & 英国人漫步大道",
     "Castle Hill (Colline du Château)\n城堡山",
     "Train: OUIGO 06:03–12:38\n（高铁）+ Walk\n（步行）",
     "Hôtel La Pérouse\nAddress: 11 Quai Rauba Capeu,\n06300 Nice\nPhone: +33 4 93 62 34 63"),
    (6, "2026/11/25\nWednesday", "France",
     "Nice → Antibes → Nice\n尼斯 → 昂蒂布 → 尼斯",
     "Regional train to Antibes (approx. 20 min)\n区域火车前往昂蒂布（约20分钟）\nOld town & coastline\n老城 & 海岸线\nOptional visit to Cannes\n可选戛纳游览\nReturn to Nice\n返回尼斯",
     "Antibes old town & port\n昂蒂布老城 & 港口\nCannes (optional)\n戛纳（可选）",
     "TER train\n（区域火车）+ Walk",
     "Novotel Suites Nice Aeroport Arenas\nAddress: 125 Boulevard René Cassin,\n06200 Nice\nPhone: +33 4 93 83 80 80"),
    (7, "2026/11/26\nThursday", "France - Italy",
     "Nice → Florence\n尼斯 → 佛罗伦萨",
     "FlixBus from Nice Airport T1 → Florence Villa Costanza\nFlixBus 尼斯机场T1 → 佛罗伦萨Villa Costanza\nTram to city center → hotel check-in\n电车入城 → 酒店入住",
     "Piazza della Signoria\n领主广场\nPonte Vecchio\n老桥\nSanta Croce district\n圣十字区",
     "FlixBus + Tram\n（长途巴士 + 电车）",
     "Hotel Loggiato Dei Serviti\nAddress: Piazza della Santissima\nAnnunziata 3, 50122 Florence\nPhone: +39 055 289 592"),
    (8, "2026/11/27\nFriday", "Italy",
     "Florence\n佛罗伦萨",
     "Florence Cathedral (morning)\n圣母百花大教堂（上午）\nLunch in Santa Croce\n圣十字区午餐\nWalk along the Arno\n沿阿诺河漫步",
     "Florence Cathedral (Duomo)\n圣母百花大教堂\nPiazza della Signoria\n领主广场\nPonte Vecchio\n老桥",
     "Walk\n（步行）",
     "B&B Cimatori\nAddress: Via Dante Alighieri 14,\n50122 Florence\nPhone: +39 055 214 213"),
    (9, "2026/11/28\nSaturday", "Italy",
     "Florence\n佛罗伦萨",
     "Uffizi Gallery (morning, pre-booked)\n乌菲兹美术馆（上午，已预订）\nLunch in city center\n市中心午餐\nBus up to Piazzale Michelangelo\n乘车前往米开朗基罗广场",
     "Uffizi Gallery\n乌菲兹美术馆\nPiazzale Michelangelo (sunset)\n米开朗基罗广场（日落）",
     "Walk / Bus\n（步行 / 公交）",
     "Hotel Indigo Florence\nAddress: Via Jacopo da Diacceto\n16/20, 50123 Florence\nPhone: +39 055 246 0400"),
    (10, "2026/11/29\nSunday", "Italy",
     "Florence → Rome\n佛罗伦萨 → 罗马",
     "Italo high-speed train → Roma Termini 10:19\nItalo高铁 → 罗马Termini 10:19\nLunch near station → Colosseum\n车站附近午餐 → 罗马斗兽场\nPre-booked visit approx. 3h\n已预订游览约3小时",
     "Colosseum\n罗马斗兽场\nPalatine Hill\n帕拉蒂尼山\nRoman Forum\n古罗马广场",
     "Train: Italo → Roma Termini 10:19\n（高铁）+ Metro",
     "Hotel in Rome\n（name & address to be confirmed）\n罗马酒店（待确认）"),
    (11, "2026/11/30\nMonday", "Italy",
     "Rome\n罗马",
     "Borghese Gallery (morning, pre-booked)\n博尔盖塞美术馆（上午，已预订）\nWalk through historic center\n历史中心步行游览",
     "Borghese Gallery\n博尔盖塞美术馆\nPantheon\n万神殿\nPiazza Navona\n纳沃纳广场\nTrevi Fountain\n特莱维喷泉\nSpanish Steps\n西班牙台阶\nPincio (sunset)\n苹丘（日落）",
     "Walk / Metro\n（步行 / 地铁）",
     "Hotel in Rome\n（name & address to be confirmed）\n罗马酒店（待确认）"),
    (12, "2026/12/01\nTuesday", "Italy - China",
     "Rome → Shanghai\n罗马 → 上海",
     "Vatican Museums & Sistine Chapel (pre-booked)\n梵蒂冈博物馆 & 西斯廷教堂（已预订）\nCastel Sant'Angelo & Ponte Sant'Angelo\n圣天使城堡 & 圣天使桥\nSt. Peter's Basilica\n圣彼得大教堂\n15:30 Leonardo Express to FCO T3\n15:30 莱昂纳多快线前往菲乌米奇诺T3",
     "Vatican Museums & Sistine Chapel\n梵蒂冈博物馆 & 西斯廷教堂\nCastel Sant'Angelo\n圣天使城堡\nSt. Peter's Basilica\n圣彼得大教堂",
     "Metro + Leonardo Express +\nFlight: CA940 19:50\n（罗马FCO T3 → 北京）",
     "On board flight\n（宿飞机上）"),
    (13, "2026/12/02\nWednesday", "China",
     "Arrive in Shanghai\n抵达上海",
     "Arrive Beijing Capital T3 13:00\n抵达北京首都T3 13:00\nConnect CA1515 (15:45)\n转机CA1515（15:45）\nArrive Shanghai Hongqiao T2 17:55\n抵达上海虹桥T2 17:55\nEnd of trip\n行程结束",
     "—",
     "Flight: CA1515 15:45–17:55\n（北京 → 上海虹桥T2）",
     "Back to home\n（返回家中）"),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Trip Itinerary"

thin = Side(style="thin")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 标题行
ws.merge_cells("A1:H1")
c = ws.cell(row=1, column=1, value="Trip Itinerary / 酒店行程单（Italy - France, 2026/11/20 – 2026/12/02）")
c.font = Font(bold=True, size=14)
c.alignment = center
ws.row_dimensions[1].height = 30

# 表头
for col, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=2, column=col, value=h)
    cell.font = Font(bold=True, size=10)
    cell.alignment = center
    cell.border = border
ws.row_dimensions[2].height = 24

# 数据行
for r, row in enumerate(ROWS, 3):
    for col, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=col, value=val)
        cell.font = Font(size=9)
        cell.alignment = center if col <= 4 else left
        cell.border = border
    ws.row_dimensions[r].height = 110

widths = [5, 13, 11, 15, 38, 34, 24, 34]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A3"
ws.print_area = f"A1:H{2 + len(ROWS)}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

out = "酒店行程单_Italy-France_2026-11.xlsx"
wb.save(out)
print("saved:", out)
